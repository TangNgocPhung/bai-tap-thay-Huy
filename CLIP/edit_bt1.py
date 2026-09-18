from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\Phung\Downloads\BT1.docx")
IMAGE_1 = Path(r"C:\Users\Phung\AppData\Local\Temp\codex-clipboard-2c2d1fdc-c0b6-4845-bf48-156bdf18c2a8.png")
IMAGE_2 = Path(r"C:\Users\Phung\AppData\Local\Temp\codex-clipboard-72bdca47-98c7-4f8f-8847-5e1cb8e2bf30.png")
OUTPUT = Path(r"C:\Users\Phung\OneDrive - Hochiminh City University of Education\Desktop\CLIP\BT1_da_kiem_tra_va_chen_hinh.docx")


def replace_paragraph_text(paragraph, new_text):
    """Replace visible text while retaining the first run's character formatting."""
    # Only inherit formatting from the actual first run.  Looking for the first
    # run that happens to have rPr can accidentally copy a later subscript,
    # bold, or language-only run onto the entire replacement paragraph.
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_rpr = deepcopy(paragraph.runs[0]._r.rPr)

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = paragraph.add_run(new_text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def find_unique_paragraph(document, starts_with):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(starts_with)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {starts_with!r}, found {len(matches)}")
    return matches[0]


def add_picture_paragraph(document, image_path, alt_text, width_inches=6.15):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    shape._inline.docPr.set("title", alt_text)
    shape._inline.docPr.set("descr", alt_text)
    return paragraph


def add_caption_paragraph(document, caption_template, text):
    paragraph = document.add_paragraph()
    if caption_template._p.pPr is not None:
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        paragraph._p.insert(0, deepcopy(caption_template._p.pPr))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)

    run = paragraph.add_run(text)
    if caption_template.runs and caption_template.runs[0]._r.rPr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(caption_template.runs[0]._r.rPr))
    else:
        run.italic = True
    return paragraph


def insert_after(anchor_paragraph, paragraphs):
    anchor = anchor_paragraph._p
    for paragraph in paragraphs:
        anchor.addnext(paragraph._p)
        anchor = paragraph._p


def main():
    document = Document(str(SOURCE))

    replacements = {
        "Thử nghiệm ban đầu là huấn luyện một CNN": (
            "Thử nghiệm ban đầu huấn luyện một CNN và một Transformer văn bản để dự đoán caption đi kèm ảnh, "
            "nhưng cách này khó mở rộng hiệu quả. Theo Hình 2 của bài báo, mô hình ngôn ngữ Transformer học khả năng "
            "zero-shot ImageNet chậm hơn khoảng 3 lần so với baseline dự đoán biểu diễn bag-of-words. Khi thay mục tiêu "
            "dự đoán bag-of-words bằng mục tiêu đối lập của CLIP, hiệu quả tăng thêm khoảng 4 lần; vì vậy mức cải thiện "
            "tổng hợp so với baseline sinh caption xấp xỉ 12 lần trong phép đo này. Với một batch gồm N cặp ảnh-văn bản, "
            "CLIP tối đa hóa cosine similarity của N cặp đúng và tối thiểu hóa similarity của N² − N cặp sai. Mô hình "
            "tối ưu symmetric cross-entropy theo cả hai chiều ảnh→văn bản và văn bản→ảnh (chi tiết ở câu 6)."
        ),
        "CLIP zero-shot đạt 76,2% top-1 accuracy": (
            "CLIP zero-shot đạt 76,2% top-1 accuracy trên ImageNet, tương đương ResNet-50 gốc được huấn luyện có giám sát "
            "trên 1,28 triệu ảnh. Cách diễn đạt chính xác là CLIP không dùng 1,28 triệu ảnh gán nhãn đó để huấn luyện riêng "
            "cho ImageNet; bài báo vẫn phân tích khả năng trùng lặp ngoài ý muốn giữa dữ liệu web tiền huấn luyện và các tập "
            "đánh giá. Trên bộ 27 dataset, CLIP zero-shot thắng baseline linear probe trên đặc trưng ResNet-50 ở 16 dataset. "
            "Prompt engineering và ensembling cải thiện gần 5 điểm phần trăm trên ImageNet; Hình 4 cũng báo cáo mức tăng gần "
            "5 điểm trung bình trên 36 dataset so với cách chỉ dùng tên lớp không có ngữ cảnh. Khi đánh giá bằng linear probe, "
            "ViT-L/14@336px vượt Noisy Student EfficientNet-L2 trên 21/27 dataset và đạt hiệu quả tính toán tốt hơn trong "
            "phạm vi so sánh của bài báo."
        ),
        "Một phát hiện quan trọng là CLIP zero-shot bền vững hơn": (
            "CLIP zero-shot bền vững hơn các mô hình có cùng độ chính xác được huấn luyện hoặc tinh chỉnh trên ImageNet khi "
            "đánh giá trên các phân phối tự nhiên khác; các mô hình CLIP thu hẹp robustness gap tới 75%. Khi thích ứng CLIP "
            "vào ImageNet bằng một bộ phân loại logistic tuyến tính có điều chuẩn L2, độ chính xác ImageNet tăng 9,2 điểm phần "
            "trăm lên 85,4%, nhưng độ chính xác trung bình dưới distribution shift lại giảm nhẹ. Kết quả này gợi ý rằng thích "
            "ứng theo một phân phối cụ thể có thể làm giảm lợi thế robustness, nhưng bài báo nhấn mạnh rằng thí nghiệm chưa đủ "
            "để kết luận quan hệ nhân quả."
        ),
        "Bài báo còn so sánh CLIP với hiệu năng của con người": (
            "Bài báo còn so sánh CLIP với con người trên Oxford-IIIT Pets ở các chế độ zero/one/two-shot và phân tích trùng "
            "lặp giữa dữ liệu tiền huấn luyện với các tập đánh giá. Tỷ lệ trùng lặp có trung vị 2,2% và trung bình 3,2%; ảnh "
            "hưởng lên độ chính xác toàn bộ tập hiếm khi vượt 0,1 điểm phần trăm. Các hạn chế được nêu gồm hiệu năng yếu trên "
            "một số tác vụ chuyên biệt, phức tạp hoặc trừu tượng như đếm vật thể, phân loại ảnh vệ tinh, nhận dạng biển báo và "
            "phát hiện khối u; CLIP cũng kém trên MNIST và few-shot logistic regression đôi khi kém hơn zero-shot."
        ),
        "Đây là phiên bản mở rộng với độ chi tiết cao hơn": (
            "Đây là phiên bản mở rộng với độ chi tiết cao hơn: cũng có 60.000 ảnh 32×32 pixel, nhưng được chia thành 100 lớp "
            "chi tiết (fine label), mỗi lớp có 600 ảnh, gồm 500 ảnh huấn luyện và 100 ảnh kiểm tra. Các lớp được nhóm thành "
            "20 superclass, mỗi superclass gồm 5 lớp. Ví dụ, superclass \"reptiles\" gồm crocodile, dinosaur, lizard, snake "
            "và turtle; superclass \"aquatic mammals\" gồm beaver, dolphin, otter, seal và whale. Mỗi ảnh vì vậy có cả nhãn "
            "fine và nhãn coarse."
        ),
        "Text encoder là một Transformer": (
            "Text encoder là một Transformer (Vaswani et al., 2017) với các điều chỉnh theo GPT-2. Cấu hình cơ sở có 63 "
            "triệu tham số, 12 lớp, độ rộng 512 và 8 attention head. Bài báo mô tả BPE viết thường với 49.152 đơn vị, còn "
            "implementation và Appendix F dùng bảng embedding 49.408 mục sau khi tính các biểu tượng byte và token đặc biệt; "
            "<|startoftext|> và <|endoftext|> có id lần lượt 49406 và 49407. Chuỗi đầu vào có context length 77, bao gồm các "
            "token đặc biệt và phần padding/truncation. Trạng thái ẩn tại <|endoftext|> được layer-normalize rồi chiếu tuyến tính "
            "vào không gian embedding chung. Text encoder dùng masked self-attention để giữ khả năng khởi tạo từ mô hình ngôn "
            "ngữ tiền huấn luyện hoặc bổ sung mục tiêu language modeling trong tương lai. Khi tăng quy mô các biến thể ResNet, "
            "nhóm tác giả tăng độ rộng nhưng giữ nguyên độ sâu của text encoder."
        ),
        "Sau khi qua encoder riêng, đặc trưng ảnh I_f": (
            "Sau khi qua encoder, đặc trưng ảnh I_f và văn bản T_f được chiếu tuyến tính bằng hai ma trận học được W_i và W_t "
            "vào cùng không gian embedding. Bài báo không dùng nonlinear projection vì không quan sát thấy lợi ích về hiệu quả "
            "huấn luyện trong thí nghiệm này. Hai embedding được chuẩn hóa L2; tích vô hướng của chúng vì thế chính là cosine "
            "similarity. Logit được nhân với exp(logit_scale), tương đương chia cho một temperature; logit_scale là tham số học "
            "được và trong quá trình huấn luyện được giới hạn để hệ số scale không vượt 100."
        ),
        "Batch size huấn luyện là 32.768": (
            "Batch size huấn luyện là 32.768, các mô hình được huấn luyện 32 epoch bằng Adam với decoupled weight decay, lịch "
            "learning rate cosine, mixed precision và các kỹ thuật tiết kiệm bộ nhớ. Với nhánh ResNet, embedding dimension "
            "không tăng đơn điệu theo kích thước mô hình: RN50, RN101, RN50x4, RN50x16 và RN50x64 lần lượt là 1024, 512, "
            "640, 768 và 1024; độ phân giải đầu vào lần lượt là 224, 224, 288, 384 và 448 pixel. Với ViT, embedding dimension "
            "là 512 cho ViT-B/32 và ViT-B/16, 768 cho ViT-L/14; độ phân giải là 224 pixel, trừ bản ViT-L/14@336px."
        ),
        "Dòng Set-ExecutionPolicy chỉ nới lỏng": (
            "Dòng Set-ExecutionPolicy chỉ nới lỏng chính sách chạy script cho cửa sổ PowerShell hiện tại và không thay đổi cấu "
            "hình toàn hệ thống. Lệnh này chỉ cần khi execution policy trên máy đang chặn script Activate.ps1; nhiều máy có thể "
            "kích hoạt môi trường ảo mà không cần bước đó."
        ),
        "Kết quả trên minh hoạ đúng bản chất": (
            "Kết quả trên minh hoạ cách dùng CLIP theo chế độ zero-shot: mô hình không được huấn luyện hoặc tinh chỉnh riêng "
            "trên CIFAR-100. Khi suy luận, danh sách 100 tên lớp được cung cấp dưới dạng prompt văn bản và các tham số mô "
            "hình không được cập nhật. Việc phân loại được thực hiện hoàn toàn thông qua so khớp trong không gian embedding chung:"
        ),
        "Tính ma trận similarity giữa tất cả các cặp": (
            "Tính ma trận similarity giữa tất cả các cặp ảnh–văn bản trong batch: logits = exp(t) · Iₑ · Tₑᵀ, trong đó "
            "t là logit_scale học được; exp(t) đóng vai trò hệ số inverse temperature."
        ),
        "Kết quả tại Hình 13 gợi ý một cơ chế nguyên nhân quan trọng": (
            "Hình 13 cho thấy mối liên hệ giữa cách huấn luyện và độ bền vững, nhưng không tự nó xác lập cơ chế nhân quả. Một "
            "giả thuyết được bài báo đặt ra là mô hình được huấn luyện hoặc thích ứng trên ImageNet có thể khai thác các tương "
            "quan giả đặc thù của phân phối đó, chẳng hạn bối cảnh, góc chụp hoặc phong cách ảnh. Các tương quan này có thể không "
            "ổn định khi nguồn dữ liệu thay đổi."
        ),
        "Ngược lại, do CLIP không được tối ưu hóa trực tiếp": (
            "Một cách giải thích khả dĩ cho lợi thế của zero-shot CLIP là mô hình không được thích ứng trực tiếp với các tập đánh "
            "giá và được tiền huấn luyện trên 400 triệu cặp ảnh-văn bản đa dạng. Tuy nhiên, bài báo lưu ý rằng lợi thế có thể đến "
            "từ nhiều yếu tố đồng thời, gồm quy mô và độ đa dạng dữ liệu, giám sát bằng ngôn ngữ và cách đánh giá zero-shot; thí "
            "nghiệm chưa tách riêng được đóng góp của từng yếu tố."
        ),
        "Một phát hiện bổ sung có liên quan mật thiết": (
            "Hình 14 bổ sung bằng thí nghiệm thích ứng CLIP vào ImageNet qua logistic regression có điều chuẩn L2 trên đặc trưng "
            "CLIP. Độ chính xác ImageNet tăng 9,2 điểm phần trăm lên 85,4%, nhưng độ chính xác trung bình trên các tập distribution "
            "shift giảm nhẹ. Hiệu năng tăng trên ImageNetV2 nhưng giảm 4,7 điểm trên ImageNet-R, 3,8 điểm trên ObjectNet, 2,8 "
            "điểm trên ImageNet-Sketch và 1,9 điểm trên ImageNet-A; thay đổi trên Youtube-BB và ImageNet-Vid không đáng kể. Các "
            "tác giả xem đây là bằng chứng gợi ý và nói rõ rằng họ chưa có kết luận chắc chắn về nguyên nhân."
        ),
        "Câu 5 đã chạy zero-shot classification": (
            "Câu 5 đã chạy zero-shot classification trên CIFAR-100 bằng repo openai/CLIP. Phần này chạy lại cùng ảnh CIFAR-100 "
            "#3637 và cùng checkpoint ViT-B/32 bằng CLIPModel/CLIPProcessor của Hugging Face. Mục tiêu là so sánh mức độ trùng "
            "khớp của kết quả top-5, độ phức tạp mã nguồn và thời gian chạy thực tế."
        ),
        "Dù thời gian đo phía transformers đã bao gồm nhiều bước hơn": (
            "Hai phép đo trên chưa có cùng phạm vi: số liệu transformers gồm tiền xử lý, tokenize và forward, còn số liệu "
            "openai/CLIP chỉ tính forward. Vì vậy không nên xem tỷ lệ 16,40/3,05 là một so sánh tốc độ mô hình được kiểm soát. "
            "Chênh lệch có thể liên quan đến phiên bản PyTorch/transformers, attention backend, warm-up, số luồng CPU và bước "
            "kiểm tra SHA-256 khi clip.load() đọc checkpoint; cần đo lại cùng đầu vào, cùng phạm vi, cùng số lần lặp và ghi rõ "
            "phiên bản thư viện trước khi kết luận nguyên nhân."
        ),
        "Về hiệu năng trong thực nghiệm này": (
            "Về hiệu năng trong lần chạy này: transformers cho thời gian quan sát thấp hơn, nhưng hai phép đo chưa cùng phạm vi "
            "nên chưa đủ để kết luận thư viện nào nhanh hơn một cách tổng quát."
        ),
        "Ba hình dưới đây minh hoạ toàn bộ luồng dữ liệu": (
            "Năm hình dưới đây minh hoạ kiến trúc và luồng dữ liệu của CLIP. Hai hình đầu cung cấp sơ đồ tổng quan; ba hình sau "
            "đối chiếu chi tiết các khối với vị trí tương ứng trong source code openai/CLIP."
        ),
        "Hình 1. Hai nhánh mã hoá": (
            "Hình 3. Hai nhánh mã hoá — Text Encoder và Vision Encoder với hai phương án ViT hoặc Modified ResNet."
        ),
        "Hình 2. Không gian embedding chung": (
            "Hình 4. Không gian embedding chung, chuẩn hoá L2, similarity nhân temperature scale và phần loss chỉ có trong bài báo."
        ),
        "Hình 3. Luồng suy luận zero-shot": (
            "Hình 5. Luồng suy luận zero-shot dùng lại hai encoder, minh hoạ bằng kết quả chạy CIFAR-100 trên ảnh số 3637."
        ),
    }

    for starts_with, new_text in replacements.items():
        paragraph = find_unique_paragraph(document, starts_with)
        replace_paragraph_text(paragraph, new_text)

    # Headings should not inherit full justification, which stretches short final lines.
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Keep TOC entries left-aligned.  The source document's justified TOC style
    # stretches short wrapped lines across the page after fields are refreshed.
    for style_name in ("TOC 1", "TOC 2", "TOC 3"):
        try:
            document.styles[style_name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        except KeyError:
            pass
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name in {"TOC 1", "TOC 2", "TOC 3"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    section_intro = find_unique_paragraph(document, "Năm hình dưới đây")
    caption_template = find_unique_paragraph(document, "Hình 3. Hai nhánh mã hoá")
    first_existing_figure = section_intro._p.getnext()
    while first_existing_figure is not None and first_existing_figure.tag != qn("w:p"):
        first_existing_figure = first_existing_figure.getnext()

    picture_1 = add_picture_paragraph(
        document,
        IMAGE_1,
        "Sơ đồ khái quát kiến trúc CLIP và suy luận zero-shot",
    )
    caption_1 = add_caption_paragraph(
        document,
        caption_template,
        "Hình 1. Sơ đồ khái quát kiến trúc CLIP, không gian embedding chung và quy trình suy luận zero-shot.",
    )
    picture_2 = add_picture_paragraph(
        document,
        IMAGE_2,
        "Sơ đồ trực quan chi tiết kiến trúc và hoạt động của CLIP",
    )
    caption_2 = add_caption_paragraph(
        document,
        caption_template,
        "Hình 2. Sơ đồ trực quan chi tiết về hai encoder, huấn luyện đối sánh và suy luận zero-shot.",
    )
    insert_after(section_intro, [picture_1, caption_1, picture_2, caption_2])

    # Keep the first pre-existing source-code figure on its own page after the two new overview diagrams.
    if first_existing_figure is not None:
        ppr = first_existing_figure.get_or_add_pPr()
        page_break_before = ppr.find(qn("w:pageBreakBefore"))
        if page_break_before is None:
            page_break_before = OxmlElement("w:pageBreakBefore")
            ppr.append(page_break_before)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
