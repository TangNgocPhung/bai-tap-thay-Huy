from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "doc_assets"
SOURCE_IMAGE = ROOT / "clip_architecture_corrected.png"
OUTPUT = ROOT / "Cau_4_Giai_thich_kien_truc_CLIP.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        r = p.add_run(item)
        set_font(r)


def add_code(doc, code):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    for idx, line in enumerate(code.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.5, color="26384A")
    p._p.get_or_add_pPr().append(parse_xml(r'<w:shd {} w:fill="F3F5F7"/>'.format(nsdecls("w"))))
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    run_prop = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    run_prop.append(style)
    math_run.append(run_prop)
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    return p


def set_picture_alt(run, alt_text):
    drawing = run._element.find(qn("w:drawing"))
    if drawing is None:
        return
    doc_prs = drawing.xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[0].set("descr", alt_text)


def add_figure(doc, path, width, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_picture_alt(run, alt_text)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_font(r, size=9, italic=True, color="4F5D6B")
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F3F7FB")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=9, color="667788")


def prepare_assets():
    ASSETS.mkdir(exist_ok=True)
    image = Image.open(SOURCE_IMAGE).convert("RGB")
    crops = {
        "clip-image-encoder.png": (20, 132, 1085, 425),
        "clip-text-encoder.png": (10, 430, 1100, 700),
        "clip-shared-space.png": (1095, 115, 1590, 700),
        "clip-zero-shot.png": (10, 690, 1590, 1035),
    }
    for filename, box in crops.items():
        image.crop(box).save(ASSETS / filename, quality=95)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    body = doc.styles["Body Text"]
    body.font.name = "Arial"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    body.font.size = Pt(11)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.15
    body.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("4F5D6B")

    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        code.font.size = Pt(8.5)
        code.paragraph_format.left_indent = Inches(0.18)
        code.paragraph_format.right_indent = Inches(0.18)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(7)
        code.paragraph_format.line_spacing = 1.0

    add_page_number(section.footer.paragraphs[0])


def build_document():
    prepare_assets()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Câu 4 Giải thích source code và kiến trúc CLIP"
    doc.core_properties.subject = "Đối chiếu các thành phần kiến trúc CLIP với source code chính thức của OpenAI"
    doc.core_properties.keywords = "CLIP, OpenAI, Image Encoder, Text Encoder, contrastive learning, zero-shot"

    # Cover and overview
    title = doc.add_paragraph(style="Title")
    title.add_run("Câu 4 Giải thích source code và kiến trúc CLIP")
    direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if direct_border is not None:
        title._p.get_or_add_pPr().remove(direct_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Đối chiếu kiến trúc với repository chính thức của OpenAI")
    set_font(r, size=13, italic=True, color="4D5D6C")

    add_body(doc, "CLIP là mô hình hai encoder. Image Encoder biến ảnh thành một vector, Text Encoder biến văn bản thành một vector có cùng số chiều, sau đó mô hình so sánh hai vector bằng cosine similarity. Điểm cốt lõi của kiến trúc là ảnh và văn bản không được ghép vào một Transformer chung; chúng được mã hóa độc lập rồi gặp nhau trong không gian embedding đa phương thức.")
    add_body(doc, "Phần trình bày dưới đây lần theo đúng luồng thực thi của repository openai/CLIP: tiền xử lý đầu vào, hai lựa chọn cho bộ mã hóa ảnh, Transformer văn bản, phép chiếu tuyến tính, chuẩn hóa L2, ma trận logits và cơ chế phân loại zero-shot. Source code phát hành chủ yếu phục vụ tải mô hình và suy luận; training loop cùng hàm loss đầy đủ được mô tả trong bài báo thay vì được đóng gói thành một hàm huấn luyện trong repository.")

    add_figure(
        doc,
        SOURCE_IMAGE,
        7.0,
        "Hình 1. Pipeline tổng thể của kiến trúc CLIP và mối liên hệ giữa huấn luyện đối sánh với suy luận zero-shot.",
        "Pipeline tổng thể CLIP gồm Image Encoder, Text Encoder, không gian embedding chung, contrastive loss và zero-shot inference.",
    )

    doc.add_page_break()

    # Section 1
    doc.add_heading("1 Tổng quan kiến trúc", level=1)
    add_body(doc, "Với một mini-batch gồm N cặp ảnh và caption tương ứng, CLIP tạo N image embeddings và N text embeddings. Sau khi chuẩn hóa, tích ma trận giữa hai tập vector tạo ra ma trận tương đồng N × N. Phần tử trên đường chéo biểu diễn cặp ảnh - văn bản đúng; các phần tử còn lại đóng vai trò cặp âm trong batch.")
    add_body(doc, "Lớp CLIP trong model.py là điểm nối của toàn bộ kiến trúc. Hàm khởi tạo tạo self.visual cho nhánh ảnh, self.transformer cho nhánh văn bản, text_projection cho phép chiếu văn bản và logit_scale cho nhiệt độ học được. Hai hàm encode_image và encode_text cung cấp embedding riêng, còn forward thực hiện chuẩn hóa và tính logits.")

    doc.add_heading("2 Tổ chức source code", level=1)
    add_table(
        doc,
        ["Tệp", "Thành phần chính", "Vai trò"],
        [
            ("clip/clip.py", "load, tokenize, _transform", "API người dùng, tải checkpoint, tiền xử lý ảnh và đóng gói token văn bản."),
            ("clip/simple_tokenizer.py", "SimpleTokenizer", "Byte-level BPE, xây dựng từ điển token và chuyển văn bản thành token ID."),
            ("clip/model.py", "ModifiedResNet, VisionTransformer, Transformer, CLIP", "Định nghĩa toàn bộ mạng neural, phép chiếu embedding và phép tính logits."),
            ("README.md", "Ví dụ zero-shot", "Minh họa cách tạo prompt lớp, mã hóa ảnh và văn bản, sau đó chọn lớp theo similarity."),
            ("model-card.md", "Mô tả mô hình", "Tóm tắt mục đích, biến thể kiến trúc, giới hạn và phạm vi sử dụng."),
        ],
        [1.3, 2.0, 3.6],
    )

    doc.add_heading("3 Tiền xử lý dữ liệu", level=1)
    doc.add_heading("3 1 Tiền xử lý ảnh", level=2)
    add_body(doc, "Hàm _transform trong clip.py đưa ảnh về đúng định dạng mà checkpoint yêu cầu. Chuỗi thao tác gồm Resize bằng nội suy bicubic, CenterCrop, chuyển sang RGB, ToTensor và Normalize bằng mean cùng standard deviation của CLIP. Kích thước đích được lấy từ model.input_resolution, vì vậy cùng một API có thể phục vụ checkpoint 224 px hoặc 336 px.")
    add_code(doc, "Resize(n_px, interpolation=BICUBIC)\nCenterCrop(n_px)\n_convert_image_to_rgb\nToTensor()\nNormalize(mean, std)")

    doc.add_heading("3 2 Tiền xử lý văn bản", level=2)
    add_body(doc, "SimpleTokenizer sử dụng byte-level BPE. Hàm clip.tokenize thêm token bắt đầu SOT và token kết thúc EOT, rồi đệm chuỗi thành context length 77. Nếu văn bản vượt giới hạn, người dùng phải cho phép truncate; khi đó token cuối vẫn được thay bằng EOT. Repository chính thức không dùng SentencePiece.")

    # Image encoder
    doc.add_heading("4 Image Encoder", level=1)
    add_body(doc, "Lớp CLIP chọn bộ mã hóa ảnh dựa trên kiểu của vision_layers. Tuple hoặc list biểu thị bốn tầng ResNet; một số nguyên biểu thị số lớp Transformer của ViT. Vì vậy cùng lớp CLIP có thể dựng các checkpoint RN50, RN101, RN50x4, ViT-B/32, ViT-B/16 hoặc ViT-L/14.")
    add_code(doc, "if isinstance(vision_layers, (tuple, list)):\n    self.visual = ModifiedResNet(...)\nelse:\n    self.visual = VisionTransformer(...)")

    doc.add_heading("4 1 Modified ResNet", level=2)
    add_body(doc, "ModifiedResNet giữ các khối Bottleneck và residual connection của ResNet nhưng thay đổi ba điểm. Phần stem dùng ba convolution 3 × 3; quá trình giảm kích thước đặt average pooling trước convolution có stride; lớp pooling cuối được thay bằng AttentionPool2d. Đây là khác biệt quan trọng so với ResNet-50 thông thường, vì CLIP không kết thúc nhánh này bằng global average pooling.")
    add_body(doc, "AttentionPool2d trải feature map NCHW thành chuỗi HW vị trí, thêm một token trung bình toàn cục và positional embedding, sau đó dùng multi-head QKV attention. Kết quả được đưa qua c_proj để có kích thước embed_dim.")

    doc.add_heading("4 2 Vision Transformer", level=2)
    add_body(doc, "VisionTransformer dùng một Conv2d có kernel_size và stride bằng patch_size để chia ảnh thành patch. Các patch được làm phẳng thành chuỗi, ghép thêm class_embedding, cộng positional_embedding và chuẩn hóa bằng ln_pre. Sau các ResidualAttentionBlock, mô hình lấy token đầu tiên, áp dụng ln_post và nhân với visual.proj. Vector trả về đã nằm trong không gian embedding chung.")

    add_figure(
        doc,
        ASSETS / "clip-image-encoder.png",
        7.0,
        "Hình 2. Hai phương án Image Encoder trong source code CLIP.",
        "Chi tiết hai nhánh Modified ResNet và Vision Transformer của Image Encoder.",
    )

    doc.add_page_break()

    # Text encoder
    doc.add_heading("5 Text Encoder", level=1)
    add_body(doc, "Text Encoder là Transformer dùng masked self-attention. Mỗi ResidualAttentionBlock thực hiện hai phép cộng residual: attention trên đầu ra LayerNorm thứ nhất và MLP trên đầu ra LayerNorm thứ hai. MLP mở rộng chiều ẩn lên bốn lần, áp dụng QuickGELU rồi chiếu về chiều ban đầu.")
    add_code(doc, "x = x + self.attention(self.ln_1(x))\nx = x + self.mlp(self.ln_2(x))")

    doc.add_heading("5 1 Causal attention mask", level=2)
    add_body(doc, "build_attention_mask tạo ma trận context_length × context_length, điền -∞ phía trên đường chéo. Khi MultiheadAttention cộng mask vào attention scores, một token không thể nhìn các token đứng sau nó. Vision Transformer không nhận mask này, nên các patch ảnh vẫn attention toàn cục.")
    add_code(doc, "mask = torch.empty(self.context_length, self.context_length)\nmask.fill_(float('-inf'))\nmask.triu_(1)")

    doc.add_heading("5 2 Tạo text embedding", level=2)
    add_body(doc, "encode_text lần lượt tạo token embedding, cộng positional embedding, đổi thứ tự tensor sang LND để phù hợp với nn.MultiheadAttention và chạy qua Transformer. Sau khi đổi lại NLD, mô hình áp dụng ln_final, lấy hidden state tại token EOT và nhân với text_projection.")
    add_code(doc, "x = self.token_embedding(text).type(self.dtype)\nx = x + self.positional_embedding.type(self.dtype)\nx = self.transformer(x.permute(1, 0, 2)).permute(1, 0, 2)\nx = self.ln_final(x).type(self.dtype)\nx = x[torch.arange(x.shape[0]), text.argmax(dim=-1)] @ self.text_projection")
    add_body(doc, "Cách dùng text.argmax(dim=-1) hoạt động vì EOT có token ID lớn nhất trong từ điển. Sau phép nhân text_projection, đầu ra có dạng [batch_size, embed_dim], cùng chiều với image embedding.")

    add_figure(
        doc,
        ASSETS / "clip-text-encoder.png",
        7.0,
        "Hình 3. Thứ tự xử lý chính xác trong Text Encoder.",
        "Text Encoder xử lý caption qua byte-level BPE, Transformer nhân quả, ln_final, EOT và text_projection.",
    )

    doc.add_page_break()

    # Shared embedding and training
    doc.add_heading("6 Không gian embedding chung", level=1)
    add_body(doc, "encode_image và encode_text đều trả về vector đã được chiếu sang embed_dim. Trong forward, hai tập vector được chuẩn hóa L2. Khi đó tích vô hướng giữa hai vector chuẩn hóa chính là cosine similarity.")
    add_equation(doc, "Î = I / ||I||₂      T̂ = T / ||T||₂")
    add_body(doc, "logit_scale được lưu ở dạng log và được lấy exp trước khi nhân với ma trận similarity. Cách tham số hóa này giữ hệ số scale dương. Giá trị khởi tạo tương ứng với temperature 0,07.")
    add_equation(doc, "S = exp(s) · Î · T̂ᵀ")
    add_code(doc, "image_features = image_features / image_features.norm(dim=1, keepdim=True)\ntext_features = text_features / text_features.norm(dim=1, keepdim=True)\nlogit_scale = self.logit_scale.exp()\nlogits_per_image = logit_scale * image_features @ text_features.t()\nlogits_per_text = logits_per_image.t()")

    doc.add_heading("7 Mục tiêu huấn luyện đối sánh", level=1)
    add_body(doc, "Với batch gồm N cặp đúng, ma trận S có kích thước N × N. Nhãn mục tiêu là các chỉ số [0, 1, ..., N - 1], nên cặp đúng của hàng i nằm ở cột i. Cross-entropy thứ nhất yêu cầu mỗi ảnh chọn đúng caption; cross-entropy thứ hai yêu cầu mỗi caption chọn đúng ảnh. Loss cuối là trung bình của hai hướng.")
    add_equation(doc, "L = ½ [CE(S, y) + CE(Sᵀ, y)]      y = (0, 1, ..., N - 1)")
    add_body(doc, "Repository openai/CLIP không định nghĩa training loop hoặc hàm contrastive loss hoàn chỉnh. Hàm forward chỉ trả logits_per_image và logits_per_text; công thức loss đối xứng xuất hiện trong bài báo. Vì vậy phần loss trong sơ đồ là thành phần của phương pháp huấn luyện, không phải một lớp có sẵn trong model.py.")

    add_figure(
        doc,
        ASSETS / "clip-shared-space.png",
        3.2,
        "Hình 4. Chuẩn hóa embedding, ma trận tương đồng và contrastive loss đối xứng.",
        "Không gian embedding chung của CLIP với chuẩn hóa L2, similarity matrix và symmetric cross-entropy.",
    )

    doc.add_page_break()

    # Zero-shot
    doc.add_heading("8 Suy luận zero shot", level=1)
    add_body(doc, "CLIP thực hiện phân loại zero-shot bằng cách biến tên lớp thành văn bản mô tả. Ví dụ, lớp dog được đặt vào prompt a photo of a dog. Text Encoder mã hóa từng prompt thành một class embedding; tập embedding này đóng vai trò bộ phân loại tuyến tính được tạo động từ ngôn ngữ.")
    add_body(doc, "Ảnh cần dự đoán được tiền xử lý và đưa vào Image Encoder. Sau khi chuẩn hóa, image embedding được nhân với ma trận class embeddings. Softmax theo chiều lớp biến các logits thành phân phối xác suất, và lớp có xác suất cao nhất được chọn.")
    add_code(doc, "text_inputs = torch.cat([\n    clip.tokenize(f'a photo of a {class_name}')\n    for class_name in classes\n])\n\nimage_features = model.encode_image(image_input)\ntext_features = model.encode_text(text_inputs)\nimage_features /= image_features.norm(dim=-1, keepdim=True)\ntext_features /= text_features.norm(dim=-1, keepdim=True)\nprobs = (100.0 * image_features @ text_features.T).softmax(dim=-1)")
    add_body(doc, "Trong luồng này, tên lớp và prompt luôn đi vào Text Encoder, còn ảnh mới luôn đi vào Image Encoder. Mô hình không cần thêm classification head cố định và cũng không cần fine-tuning nếu chỉ thực hiện zero-shot inference.")

    add_figure(
        doc,
        ASSETS / "clip-zero-shot.png",
        7.0,
        "Hình 5. Pipeline phân loại zero-shot của CLIP.",
        "Hai nhánh zero-shot gồm prompt vào Text Encoder và ảnh mới vào Image Encoder, sau đó cosine similarity và softmax.",
    )

    doc.add_heading("9 Tải checkpoint và khôi phục kiến trúc", level=1)
    add_body(doc, "Hàm clip.load tải checkpoint theo tên mô hình hoặc nhận đường dẫn cục bộ. Khi dùng state_dict, build_model suy ra cấu hình kiến trúc từ hình dạng của trọng số. Nếu state_dict chứa visual.proj thì checkpoint dùng ViT; nếu không, hàm đếm số Bottleneck trong visual.layer1 đến visual.layer4 để dựng Modified ResNet. Cách này giúp repository không phải hard-code toàn bộ tham số của từng biến thể.")
    add_body(doc, "Sau khi dựng mô hình, convert_weights chuyển các convolution, linear, projection và tham số attention thích hợp sang FP16. State dictionary được nạp và mô hình được trả về ở chế độ eval. clip.load đồng thời trả preprocess tương ứng với input_resolution của checkpoint.")

    doc.add_page_break()

    # Mapping table and conclusion
    doc.add_heading("10 Đối chiếu kiến trúc với source code", level=1)
    add_table(
        doc,
        ["Thành phần kiến trúc", "Lớp hoặc hàm", "Kết quả"],
        [
            ("Tiền xử lý ảnh", "clip.py::_transform", "Tensor ảnh [N, 3, H, W] đã resize, crop và normalize."),
            ("Tokenization", "SimpleTokenizer; clip.tokenize", "Tensor token [N, 77] có SOT và EOT."),
            ("ResNet image encoder", "Bottleneck; ModifiedResNet; AttentionPool2d", "Image embedding sau c_proj."),
            ("ViT image encoder", "VisionTransformer", "Image embedding lấy từ class token sau visual.proj."),
            ("Khối Transformer", "ResidualAttentionBlock; Transformer", "Chuỗi đặc trưng sau self-attention và MLP."),
            ("Text encoder", "CLIP.encode_text", "Vector EOT sau ln_final và text_projection."),
            ("Embedding ảnh", "CLIP.encode_image", "Đầu ra self.visual có kích thước embed_dim."),
            ("Similarity logits", "CLIP.forward", "logits_per_image và logits_per_text."),
            ("Khôi phục mô hình", "build_model", "Kiến trúc được suy ra từ state_dict."),
            ("Zero-shot", "README.md example", "Xác suất trên các lớp được tạo từ prompt văn bản."),
        ],
        [1.75, 2.25, 2.9],
    )

    doc.add_heading("11 Kết luận", level=1)
    add_body(doc, "Source code phản ánh trực tiếp thiết kế dual encoder của CLIP. Nhánh ảnh dùng Modified ResNet có AttentionPool2d hoặc Vision Transformer; nhánh văn bản dùng causal Transformer và lấy biểu diễn tại EOT. Hai đầu ra được chiếu sang cùng không gian, chuẩn hóa và so sánh bằng cosine similarity có scale học được. Cùng cơ chế này hỗ trợ contrastive pre-training và cho phép xây dựng bộ phân loại zero-shot từ tên lớp bằng ngôn ngữ tự nhiên.")

    doc.add_heading("Tài liệu tham khảo", level=1)
    references = [
        ("OpenAI CLIP repository", "https://github.com/openai/CLIP"),
        ("Source code model.py", "https://github.com/openai/CLIP/blob/main/clip/model.py"),
        ("Source code clip.py", "https://github.com/openai/CLIP/blob/main/clip/clip.py"),
        ("Radford và cộng sự Learning Transferable Visual Models From Natural Language Supervision", "https://arxiv.org/abs/2103.00020"),
    ]
    for index, (label, url) in enumerate(references, start=1):
        p = doc.add_paragraph(style="Body Text")
        r = p.add_run(f"{index}. {label}. ")
        set_font(r)
        add_hyperlink(p, url, url)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
